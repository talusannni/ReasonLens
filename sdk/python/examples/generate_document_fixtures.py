"""Generate real PDF and DOCX fixtures for ReasonLens document-RAG tests."""
from pathlib import Path
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent / "data" / "documents"
ROOT.mkdir(parents=True, exist_ok=True)

def make_pdf() -> Path:
    path = ROOT / "travel_policy_2026.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    _, height = A4
    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, height - 72, "Travel Policy 2026")
    c.setFont("Helvetica", 11)
    lines = [
        "International hotel expenses are reimbursable up to INR 12000 per night.",
        "Claims require an itemized hotel invoice and manager approval.",
        "Airfare booked through an approved travel agency is reimbursable.",
        "A cancelled flight may be claimed only when the cancellation was",
        "caused by the airline or another documented operational disruption.",
    ]
    y = height - 120
    for line in lines:
        c.drawString(72, y, line)
        y -= 20
    c.showPage()
    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, height - 72, "Travel Policy 2026 - Exceptions")
    c.setFont("Helvetica", 11)
    c.drawString(72, height - 120, "Personal upgrades and minibar expenses are not reimbursable.")
    c.save()
    return path

def make_docx() -> Path:
    path = ROOT / "expense_policy.docx"
    doc = Document()
    doc.add_heading("Expense Reimbursement Policy", level=1)
    doc.add_paragraph("Domestic travel: employees may claim up to INR 5000 per hotel night.")
    doc.add_paragraph("International travel: employees may claim up to INR 12000 per hotel night.")
    doc.add_paragraph("Required evidence: an itemized invoice and manager approval.")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ["Expense", "International Limit", "Approval"]):
        cell.text = value
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text = "Hotel", "INR 12000/night", "Manager"
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text = "Minibar", "Not reimbursable", "N/A"
    doc.save(path)
    return path

if __name__ == "__main__":
    print(f"PDF: {make_pdf()}")
    print(f"DOCX: {make_docx()}")
